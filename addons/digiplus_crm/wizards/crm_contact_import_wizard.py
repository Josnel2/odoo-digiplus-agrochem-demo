import base64
import csv
import io

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from ..models.selections import (
    BUSINESS_SECTOR_SELECTION,
    CONTACT_LANGUAGE_SELECTION,
    PRIORITY_SELECTION,
    SERVICE_SELECTION,
)

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - runtime dependency in container
    load_workbook = None


class CrmContactImportWizard(models.TransientModel):
    _name = "digiplus.crm.contact.import.wizard"
    _description = "Import Contacts CRM DigiPlus"

    data_file = fields.Binary(string="Fichier", required=True)
    filename = fields.Char(string="Nom du fichier")
    line_ids = fields.One2many(
        "digiplus.crm.contact.import.line",
        "wizard_id",
        string="Lignes analysees",
    )
    create_opportunities = fields.Boolean(
        string="Creer une opportunite si un service est renseigne",
        default=True,
    )
    summary = fields.Text(string="Resume", readonly=True)
    line_count = fields.Integer(string="Lignes", compute="_compute_counts")
    create_count = fields.Integer(string="A creer", compute="_compute_counts")
    link_count = fields.Integer(string="A lier", compute="_compute_counts")
    validate_count = fields.Integer(string="A valider", compute="_compute_counts")
    reject_count = fields.Integer(string="Rejetees", compute="_compute_counts")

    @api.depends("line_ids.action_recommended")
    def _compute_counts(self):
        for wizard in self:
            wizard.line_count = len(wizard.line_ids)
            wizard.create_count = len(wizard.line_ids.filtered(lambda line: line.action_recommended == "create"))
            wizard.link_count = len(wizard.line_ids.filtered(lambda line: line.action_recommended == "link"))
            wizard.validate_count = len(
                wizard.line_ids.filtered(lambda line: line.action_recommended == "validate")
            )
            wizard.reject_count = len(wizard.line_ids.filtered(lambda line: line.action_recommended == "reject"))

    def _decode_file(self):
        self.ensure_one()
        if not self.data_file:
            raise UserError(_("Veuillez charger un fichier CSV ou Excel."))
        return base64.b64decode(self.data_file)

    def _get_rows_from_file(self):
        self.ensure_one()
        raw_data = self._decode_file()
        filename = (self.filename or "").lower()
        if filename.endswith(".xlsx") or filename.endswith(".xlsm"):
            if load_workbook is None:
                raise UserError(_("Le support Excel requiert la librairie openpyxl dans l'environnement Odoo."))
            workbook = load_workbook(io.BytesIO(raw_data), read_only=True, data_only=True)
            sheet = workbook.active
            headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            return [
                {headers[index]: str(cell.value or "").strip() for index, cell in enumerate(row)}
                for row in sheet.iter_rows(min_row=2)
                if any(cell.value not in (None, "") for cell in row)
            ]
        text_stream = io.StringIO(raw_data.decode("utf-8-sig"))
        reader = csv.DictReader(text_stream)
        return [{key.strip(): (value or "").strip() for key, value in row.items()} for row in reader]

    def _normalize_selection_value(self, value, selection):
        normalized = (value or "").strip().lower()
        if not normalized:
            return False
        labels = {label.lower(): key for key, label in selection}
        keys = {key.lower(): key for key, _label in selection}
        return labels.get(normalized) or keys.get(normalized)

    def _find_duplicates(self, company_name, contact_name, email, phone, website):
        Partner = self.env["res.partner"]
        normalized_email = Partner.normalize_email(email)
        normalized_phone = Partner.normalize_phone(phone)

        exact_match = self.env["res.partner"]
        if normalized_email:
            for partner in Partner.search([("email", "!=", False)]):
                if Partner.normalize_email(partner.email) == normalized_email:
                    exact_match |= partner
        if len(exact_match) == 1:
            return "exact", exact_match
        if len(exact_match) > 1:
            return "ambiguous", exact_match

        reliable = self.env["res.partner"]
        if company_name and normalized_phone:
            candidates = Partner.search([("name", "ilike", company_name)])
            for partner in candidates:
                if Partner.normalize_phone(partner.phone) == normalized_phone:
                    reliable |= partner
        if len(reliable) == 1:
            return "reliable", reliable
        if len(reliable) > 1:
            return "ambiguous", reliable

        ambiguous = self.env["res.partner"]
        if company_name:
            ambiguous |= Partner.search([("name", "ilike", company_name)])
        if website:
            ambiguous |= Partner.search([("website", "ilike", website)])
        if contact_name:
            ambiguous |= Partner.search([("name", "ilike", contact_name)])
        if len(ambiguous) == 1:
            return "ambiguous", ambiguous
        if len(ambiguous) > 1:
            return "ambiguous", ambiguous
        return False, self.env["res.partner"]

    def action_parse_file(self):
        self.ensure_one()
        rows = self._get_rows_from_file()
        self.line_ids.unlink()
        line_values = []
        required_headers = {"company_name", "contact_name", "email", "phone", "service_requested", "language"}

        for index, row in enumerate(rows, start=1):
            missing_headers = [header for header in required_headers if header not in row]
            if missing_headers:
                raise UserError(
                    _("Colonnes manquantes dans le fichier: %s.") % ", ".join(sorted(missing_headers))
                )
            service_key = self._normalize_selection_value(row.get("service_requested"), SERVICE_SELECTION)
            sector_key = self._normalize_selection_value(row.get("sector"), BUSINESS_SECTOR_SELECTION)
            priority_key_raw = self._normalize_selection_value(row.get("priority"), PRIORITY_SELECTION)
            language_key_raw = self._normalize_selection_value(row.get("language"), CONTACT_LANGUAGE_SELECTION)
            priority_key = priority_key_raw or "medium"
            language_key = language_key_raw or "fr_FR"
            duplicate_state, partners = self._find_duplicates(
                row.get("company_name"),
                row.get("contact_name"),
                row.get("email"),
                row.get("phone") or row.get("mobile"),
                row.get("website"),
            )

            action_recommended = "create"
            duplicate_notes = False
            errors = []

            if not row.get("company_name") and not row.get("contact_name"):
                errors.append(_("Societe ou contact obligatoire"))
            if not row.get("email") and not (row.get("phone") or row.get("mobile")):
                errors.append(_("Email ou telephone obligatoire"))
            if row.get("service_requested") and not service_key:
                errors.append(_("Service non reconnu"))
            if row.get("sector") and not sector_key:
                errors.append(_("Secteur non reconnu"))
            if row.get("priority") and not priority_key_raw:
                errors.append(_("Priorite non reconnue"))
            if row.get("language") and not language_key_raw:
                errors.append(_("Langue non reconnue"))

            if duplicate_state == "exact":
                action_recommended = "link"
                duplicate_notes = _("Doublon exact detecte par email")
            elif duplicate_state == "reliable":
                action_recommended = "link"
                duplicate_notes = _("Doublon fiable detecte par societe + telephone")
            elif duplicate_state == "ambiguous":
                action_recommended = "validate"
                duplicate_notes = _("Doublon ambigu a valider manuellement")

            if errors:
                action_recommended = "reject"

            line_values.append(
                {
                    "wizard_id": self.id,
                    "sequence": index,
                    "company_name": row.get("company_name"),
                    "contact_name": row.get("contact_name"),
                    "email": row.get("email"),
                    "phone": row.get("phone"),
                    "mobile": row.get("mobile"),
                    "job_title": row.get("job_title"),
                    "website": row.get("website"),
                    "sector": sector_key,
                    "service_requested": service_key,
                    "source_name": row.get("source"),
                    "salesperson_name": row.get("salesperson"),
                    "priority": priority_key,
                    "notes": row.get("notes"),
                    "language": language_key,
                    "duplicate_state": duplicate_state or "none",
                    "action_recommended": action_recommended,
                    "existing_partner_id": partners[:1].id if partners else False,
                    "duplicate_notes": duplicate_notes,
                    "error_message": "\n".join(errors) if errors else False,
                }
            )

        self.env["digiplus.crm.contact.import.line"].create(line_values)
        self.summary = _(
            "%s lignes analysees, %s a creer, %s a lier, %s a valider, %s rejetees."
        ) % (
            len(line_values),
            len([line for line in line_values if line["action_recommended"] == "create"]),
            len([line for line in line_values if line["action_recommended"] == "link"]),
            len([line for line in line_values if line["action_recommended"] == "validate"]),
            len([line for line in line_values if line["action_recommended"] == "reject"]),
        )
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "view_mode": "form",
            "res_id": self.id,
            "target": "new",
        }

    def _resolve_salesperson(self, salesperson_name):
        if not salesperson_name:
            return self.env.user
        user = self.env["res.users"].search([("name", "ilike", salesperson_name)], limit=1)
        return user or self.env.user

    def _resolve_source(self, source_name):
        if not source_name:
            return False
        source = self.env["utm.source"].search([("name", "ilike", source_name)], limit=1)
        return source

    def action_apply_import(self):
        self.ensure_one()
        if not self.line_ids:
            raise UserError(_("Aucune ligne n'a ete analysee."))

        created_partners = self.env["res.partner"]
        created_leads = self.env["crm.lead"]

        for line in self.line_ids:
            if line.action_recommended in ("reject", "validate"):
                continue
            partner = line.existing_partner_id
            if line.action_recommended == "create" or not partner:
                partner = self.env["res.partner"].create(line.prepare_partner_vals())
                created_partners |= partner

            if self.create_opportunities and line.service_requested:
                lead = self.env["crm.lead"].create(line.prepare_lead_vals(partner))
                created_leads |= lead

        message = _("%s partenaires crees, %s opportunites creees.") % (len(created_partners), len(created_leads))
        self.summary = "%s\n%s" % ((self.summary or ""), message)
        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Import CRM termine"),
                "message": message,
                "type": "success",
                "sticky": False,
            },
        }


class CrmContactImportLine(models.TransientModel):
    _name = "digiplus.crm.contact.import.line"
    _description = "Ligne d'import CRM DigiPlus"
    _order = "sequence"

    wizard_id = fields.Many2one("digiplus.crm.contact.import.wizard", required=True, ondelete="cascade")
    sequence = fields.Integer()
    company_name = fields.Char(string="Societe")
    contact_name = fields.Char(string="Contact")
    email = fields.Char(string="Email")
    phone = fields.Char(string="Telephone")
    mobile = fields.Char(string="Mobile")
    job_title = fields.Char(string="Fonction")
    website = fields.Char(string="Site web")
    sector = fields.Selection(BUSINESS_SECTOR_SELECTION, string="Secteur")
    service_requested = fields.Selection(SERVICE_SELECTION, string="Service")
    source_name = fields.Char(string="Source")
    salesperson_name = fields.Char(string="Commercial")
    priority = fields.Selection(PRIORITY_SELECTION, string="Priorite")
    notes = fields.Text(string="Notes")
    language = fields.Selection(CONTACT_LANGUAGE_SELECTION, string="Langue")
    duplicate_state = fields.Selection(
        [
            ("none", "Aucun"),
            ("exact", "Exact"),
            ("reliable", "Fiable"),
            ("ambiguous", "Ambigu"),
        ],
        string="Doublon",
    )
    action_recommended = fields.Selection(
        [
            ("create", "Creer"),
            ("link", "Lier a l'existant"),
            ("validate", "A valider"),
            ("reject", "Rejeter"),
        ],
        string="Action",
    )
    existing_partner_id = fields.Many2one("res.partner", string="Partenaire existant")
    duplicate_notes = fields.Char(string="Analyse doublon")
    error_message = fields.Text(string="Erreur")

    def prepare_partner_vals(self):
        self.ensure_one()
        partner_lang = self.env["res.partner"].sanitize_installed_lang(self.language)
        return {
            "name": self.company_name or self.contact_name,
            "company_type": "company",
            "email": self.email,
            "phone": self.phone,
            "mobile": self.mobile,
            "website": self.website,
            "function": self.job_title,
            "lang": partner_lang,
            "x_contact_language": self.language,
            "x_partner_business_sector": self.sector,
            "x_partner_priority_level": self.priority,
            "x_preferred_service_requested": self.service_requested,
            "x_partner_account_notes": self.notes,
        }

    def prepare_lead_vals(self, partner):
        self.ensure_one()
        source = self.wizard_id._resolve_source(self.source_name)
        salesperson = self.wizard_id._resolve_salesperson(self.salesperson_name)
        stage = self.env.ref("digiplus_crm.stage_prospect", raise_if_not_found=False)
        return {
            "name": "%s - %s" % (
                self.company_name or partner.name,
                dict(self._fields["service_requested"].selection).get(self.service_requested, _("Besoin commercial")),
            ),
            "type": "opportunity",
            "stage_id": stage.id if stage else False,
            "partner_id": partner.id,
            "partner_name": self.company_name or partner.name,
            "contact_name": self.contact_name,
            "email_from": self.email,
            "phone": self.phone or self.mobile,
            "description": self.notes,
            "source_id": source.id if source else False,
            "user_id": salesperson.id,
            "x_service_requested": self.service_requested,
            "x_business_sector": self.sector,
            "x_priority_level": self.priority,
            "x_contact_language": self.language,
            "x_need_type": "implementation",
        }
